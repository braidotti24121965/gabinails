from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = "Relatorio_de_Desempenho_das_Campanhas_de_Email.docx"

NAVY = "263247"
ROSE = "A95670"
PALE = "F7EEF1"
PALE_BLUE = "EEF2F7"
LIGHT = "D9D9D9"
MID = "667085"
BLACK = "000000"
WHITE = "FFFFFF"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=LIGHT, size="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = tcBorders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_cell_text(cell, text, bold=False, color=BLACK, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)
    borders(cell)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(hdr.cells[i], NAVY)
        if widths:
            hdr.cells[i].width = widths[i]
    for ridx, row in enumerate(rows):
        new_row = table.add_row()
        prevent_row_split(new_row)
        cells = new_row.cells
        for i, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[i], val, color=BLACK, size=10.2, align=align)
            shade(cells[i], WHITE if ridx % 2 == 0 else PALE_BLUE)
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, lead, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(lead)
    r.bold = True
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.8)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.13

for style_name, size, before, after in (("Title", 27, 0, 10), ("Heading 1", 16, 12, 6), ("Heading 2", 12.5, 9, 4)):
    s = styles[style_name]
    s.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(BLACK)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

# Neutralize the decorative bottom rule carried by some built-in Title styles.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_pbdr = title_ppr.find(qn("w:pBdr"))
if title_pbdr is not None:
    title_ppr.remove(title_pbdr)
title_pbdr = OxmlElement("w:pBdr")
for edge in ("top", "left", "bottom", "right", "between"):
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "nil")
    title_pbdr.append(el)
title_ppr.append(title_pbdr)

styles["List Bullet"].font.name = "Aptos"
styles["List Bullet"].font.size = Pt(10.6)
styles["List Bullet"].paragraph_format.space_after = Pt(4)
styles["List Bullet"].paragraph_format.left_indent = Inches(0.22)
styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.16)

# Header and footer
header = section.header
hp = header.paragraphs[0]
hp.text = "RELATÓRIO DE CAMPANHAS DE E MAIL"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.runs[0]
hr.font.name = "Aptos"
hr.font.size = Pt(8)
hr.font.bold = True
hr.font.color.rgb = RGBColor.from_string(MID)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Período analisado  21 de agosto a 21 de setembro de 2026   •   ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
fp._p.append(fld)
for r in fp.runs:
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID)

# Cover / opening
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(30)
r = p.add_run("ANÁLISE DE DESEMPENHO")
r.font.name = "Aptos"
r.font.size = Pt(10)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(ROSE)

t = doc.add_paragraph(style="Title")
t.add_run("Relatório de desempenho das campanhas de email")
# Remove any border inherited from the built-in title style.
pPr = t._p.get_or_add_pPr()
pBdr = pPr.find(qn("w:pBdr"))
if pBdr is not None:
    pPr.remove(pBdr)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(22)
rr = sub.add_run("21 campanhas enviadas entre 21 de agosto e 21 de setembro de 2026")
rr.font.name = "Aptos"
rr.font.size = Pt(13)
rr.font.color.rgb = RGBColor.from_string(MID)

add_heading(doc, "Visão executiva", 1)
add_body(doc, "No período analisado, foram realizadas 21 campanhas para 4.155 destinatários. Após 999 rejeições de entrega, estima-se que 3.156 mensagens tenham sido entregues. Essas entregas geraram 1.484 aberturas, 80 cliques e 12 cancelamentos de assinatura.")
add_body(doc, "A taxa de abertura de 47,02% indica boa capacidade de gerar interesse entre as mensagens que chegaram aos destinatários. O principal ponto de atenção está antes dessa etapa: 24,04% dos endereços resultaram em soft ou hard bounce, reduzindo a entrega estimada a 75,96%. Assim, a oportunidade mais imediata é melhorar a qualidade da base e a entregabilidade, preservando o desempenho de abertura e ampliando o volume efetivamente alcançado.")

add_heading(doc, "Indicadores principais", 1)
summary = doc.add_table(rows=2, cols=3)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
cards = [
    ("21", "campanhas enviadas"),
    ("3.156", "entregas estimadas"),
    ("47,02%", "taxa de abertura"),
    ("24,04%", "taxa de rejeição"),
    ("2,53%", "taxa de cliques"),
    ("0,38%", "cancelamento de assinatura"),
]
for idx, (value, label) in enumerate(cards):
    cell = summary.rows[idx // 3].cells[idx % 3]
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    rv = p1.add_run(value)
    rv.font.name = "Aptos Display"
    rv.font.size = Pt(19)
    rv.font.bold = True
    rv.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    rl = p2.add_run(label)
    rl.font.name = "Aptos"
    rl.font.size = Pt(9.3)
    rl.font.color.rgb = RGBColor.from_string(MID)
    shade(cell, PALE if idx in (2, 5) else WHITE)
    borders(cell)
    margins(cell, top=180, start=120, bottom=180, end=120)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(10)
note.paragraph_format.space_after = Pt(0)
rn = note.add_run("Conclusão central  ")
rn.bold = True
rn.font.color.rgb = RGBColor.from_string(ROSE)
note.add_run("há interesse após a entrega, mas a redução das rejeições é necessária para converter esse interesse em maior alcance.")

doc.add_page_break()

add_heading(doc, "Resultados consolidados", 1)
add_body(doc, "A tabela apresenta os dados observados na plataforma e os indicadores calculados a partir deles. As taxas de abertura, clique e cancelamento exibidas na imagem usam como base as 3.156 entregas estimadas.")
add_table(
    doc,
    ["Indicador", "Resultado", "Base de cálculo"],
    [
        ("Campanhas enviadas", "21", "Dado da plataforma"),
        ("Destinatários", "4.155", "Dado da plataforma"),
        ("Soft e hard bounces", "999", "Dado da plataforma"),
        ("Entregas estimadas", "3.156", "4.155 − 999"),
        ("Aberturas", "1.484", "Dado da plataforma"),
        ("Taxa de abertura", "47,02%", "1.484 ÷ 3.156"),
        ("Cliques", "80", "Dado da plataforma"),
        ("Taxa de cliques", "2,53%", "80 ÷ 3.156"),
        ("Respondidos", "0", "Dado da plataforma"),
        ("Cancelamentos", "12", "Dado da plataforma"),
        ("Taxa de cancelamento", "0,38%", "12 ÷ 3.156"),
    ],
    [Inches(2.6), Inches(1.25), Inches(2.7)],
)

add_heading(doc, "Funil de desempenho", 1)
add_body(doc, "O funil mostra onde o volume é reduzido ao longo da jornada. As aberturas e os cliques são registros agregados informados pela plataforma; a imagem não permite confirmar se são eventos únicos por pessoa.")
add_table(
    doc,
    ["Etapa", "Volume", "% dos destinatários", "% da etapa anterior"],
    [
        ("Destinatários", "4.155", "100,00%", "—"),
        ("Entregas estimadas", "3.156", "75,96%", "75,96%"),
        ("Aberturas", "1.484", "35,72%", "47,02%"),
        ("Cliques", "80", "1,93%", "5,39%"),
        ("Cancelamentos", "12", "0,29%", "0,81% das aberturas"),
    ],
    [Inches(1.85), Inches(1.05), Inches(1.65), Inches(2.0)],
)

doc.add_page_break()
add_heading(doc, "Médias por campanha", 1)
add_body(doc, "Como a imagem apresenta apenas totais consolidados, as médias abaixo dividem cada resultado pelas 21 campanhas e servem como referência de escala, não como avaliação de campanhas individuais.")
add_table(
    doc,
    ["Métrica", "Média por campanha"],
    [
        ("Destinatários", "197,86"),
        ("Entregas estimadas", "150,29"),
        ("Rejeições", "47,57"),
        ("Aberturas", "70,67"),
        ("Cliques", "3,81"),
        ("Cancelamentos", "0,57"),
    ],
    [Inches(3.6), Inches(2.95)],
)

add_heading(doc, "Leitura dos resultados", 1)
add_heading(doc, "Entregabilidade", 2)
add_body(doc, "Foram registradas 999 rejeições entre 4.155 destinatários, o equivalente a 24,04%. Com isso, aproximadamente um em cada quatro endereços não recebeu a mensagem. Como o dado reúne soft e hard bounces, não é possível identificar pela imagem quanto decorreu de falhas temporárias e quanto decorreu de endereços inválidos ou indisponíveis de forma permanente.")

add_heading(doc, "Abertura e interesse inicial", 2)
add_body(doc, "Entre as entregas estimadas, houve 1.484 aberturas, correspondentes a 47,02%. Esse resultado mostra que assunto, remetente e momento de envio conseguiram gerar abertura em parcela relevante das mensagens entregues. A plataforma informa que aberturas e cliques automatizados estão incluídos; portanto, esses eventos podem conter atividade de filtros de segurança ou mecanismos de privacidade.")

add_heading(doc, "Cliques e avanço para a ação", 2)
add_body(doc, "Os 80 cliques representam 2,53% das entregas e 5,39% das aberturas. A diferença entre abertura e clique indica espaço para fortalecer a ligação entre a promessa do assunto, o conteúdo da mensagem, a oferta e a chamada para ação.")

add_heading(doc, "Respostas e cancelamentos", 2)
add_body(doc, "Não houve respostas registradas. Esse dado deve ser interpretado de acordo com o objetivo das campanhas: se a meta era direcionar tráfego por links, a ausência de respostas pode ser esperada; se havia expectativa de conversa por email, convém revisar o convite à resposta e a configuração do endereço remetente. Os 12 cancelamentos equivalem a 0,38% das entregas.")

doc.add_page_break()
add_heading(doc, "Recomendações para o próximo ciclo", 1)
add_bullet(doc, "Priorizar a higiene da base. ", "Separar soft e hard bounces, suprimir endereços com rejeição permanente e revisar as origens de captação que concentram falhas.")
add_bullet(doc, "Proteger a reputação de envio. ", "Evitar novos disparos para endereços inválidos, acompanhar autenticações do domínio e observar a evolução da taxa de entrega em cada campanha.")
add_bullet(doc, "Aumentar a conversão de abertura em clique. ", "Testar uma chamada principal mais evidente, reduzir a dispersão de links e alinhar o conteúdo da mensagem à promessa feita no assunto.")
add_bullet(doc, "Comparar campanhas individualmente. ", "Analisar assunto, público, data, horário e oferta de cada envio para identificar quais combinações explicam as melhores aberturas e os melhores cliques.")
add_bullet(doc, "Definir o papel da resposta. ", "Se respostas forem desejadas, incluir uma pergunta objetiva e confirmar que o endereço de resposta é monitorado; caso contrário, tratar cliques e conversões no destino como métricas principais.")
add_bullet(doc, "Medir além do clique. ", "Adicionar acompanhamento de conversões no site ou canal de destino para conectar os envios a agendamentos, solicitações ou vendas.")

add_heading(doc, "Plano de acompanhamento", 1)
add_table(
    doc,
    ["Prioridade", "Ação", "Indicador de acompanhamento"],
    [
        ("Imediata", "Limpar e segmentar a base", "Bounces e taxa de entrega"),
        ("Próximos envios", "Testar assunto, CTA e conteúdo", "Abertura, cliques e clique por abertura"),
        ("Próximo ciclo", "Rastrear conversões após o clique", "Agendamentos, leads ou vendas"),
    ],
    [Inches(1.2), Inches(3.05), Inches(2.3)],
)

add_heading(doc, "Metodologia e limitações", 1)
add_body(doc, "Fonte: captura da tela de Estatísticas referente ao período de 21/08/2026 a 21/09/2026. Entregas estimadas = destinatários − soft e hard bounces. Taxa de rejeição = bounces ÷ destinatários. Clique por abertura = cliques ÷ aberturas. Médias por campanha = total ÷ 21.")
add_body(doc, "A captura não contém resultados separados por campanha, distinção entre soft e hard bounce, identificação de eventos únicos, conversões após o clique ou comparativos históricos. Por isso, o relatório não atribui causas específicas nem compara o desempenho com benchmarks externos. A própria plataforma informa que aberturas e cliques automatizados estão incluídos.")

# Core properties
doc.core_properties.title = "Relatório de desempenho das campanhas de email"
doc.core_properties.subject = "Resultados consolidados de 21 campanhas de email"
doc.core_properties.author = "Equipe de atendimento"
doc.core_properties.comments = "Elaborado a partir da captura de estatísticas fornecida."

doc.save(OUT)
print(OUT)
